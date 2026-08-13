from apify_client import ApifyClient
from openpyxl import Workbook
from config import APIFY_TOKEN
import os

client = ApifyClient(APIFY_TOKEN)

categories = {
    "Veri Analisti": "Data Analyst",
    "İş Analisti": "Business Analyst",
    "BI_Raporlama": "BI Reporting",
    "Yazılım Geliştirici": "Software Developer"
}

all_items = []

for category, keyword in categories.items():

    print(f"\n--- {category} ilanları çekiliyor ---")

    url = (
        "https://www.linkedin.com/jobs/search/"
        f"?keywords={keyword.replace(' ', '%20')}"
        "&location=Turkey"
    )

    run_input = {
        "urls": [url],
        "count": 20
    }

    run = client.actor(
        "curious_coder/linkedin-jobs-scraper"
    ).call(
        run_input=run_input
    )

    items = client.dataset(
        run.default_dataset_id
    ).list_items().items

    if items:
      print("\nİLK İLANIN ALANLARI:")
      print(items[0].keys())
      print("\nİLK İLAN:")
      print(items[0])

    print(f"{len(items)} ilan bulundu.")

    for item in items:

        all_items.append({
     "Kategori": category,
     "Pozisyon": item.get("title"),
     "Şirket": item.get("companyName"),
     "Şehir": item.get("location"),
     "İlan Tarihi": item.get("postedAt"),
     "Çalışma Tipi": item.get("employmentType"),
     "Deneyim Seviyesi": item.get("seniorityLevel"),
     "Maaş": item.get("salary"),
     "Başvuran Sayısı": item.get("applicantsCount"),
     "Açıklama": item.get("descriptionText"),
     "Link": item.get("link")
    })


# --------------------------------------------------
# TEKRAR EDEN İLANLARI ENGELLE
# --------------------------------------------------

unique_items = {}
    
for item in all_items:

    link = item["Link"]

    if link:
        unique_items[link] = item


all_items = list(unique_items.values())


# --------------------------------------------------
# EXCEL OLUŞTUR
# --------------------------------------------------

output_folder = "04_Python/output"

os.makedirs(output_folder, exist_ok=True)

file_path = os.path.join(
    output_folder,
    "is_ilanlari_ham.xlsx"
)

wb = Workbook()
ws = wb.active
ws.title = "Ham_Veriler"

headers = [
    "Kategori",
    "Pozisyon",
    "Şirket",
    "Şehir",
    "İlan Tarihi",
    "Çalışma Tipi",
    "Deneyim Seviyesi",
    "Maaş",
    "Başvuran Sayısı",
    "Açıklama",
    "Link"
]

ws.append(headers)


for item in all_items:

    ws.append([
    item["Kategori"],
    item["Pozisyon"],
    item["Şirket"],
    item["Şehir"],
    item["İlan Tarihi"],
    item["Çalışma Tipi"],
    item["Deneyim Seviyesi"],
    item["Maaş"],
    item["Başvuran Sayısı"],
    item["Açıklama"],
    item["Link"]
])


# --------------------------------------------------
# TABLO BİÇİMİ
# --------------------------------------------------

from openpyxl.worksheet.table import Table, TableStyleInfo

last_row = ws.max_row

table = Table(
    displayName="IsIlanlari",
    ref=f"A1:K{last_row}"
)

style = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False
)

table.tableStyleInfo = style

ws.add_table(table)

ws.freeze_panes = "A2"

# Sütun genişlikleri
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 35
ws.column_dimensions["C"].width = 30
ws.column_dimensions["D"].width = 25
ws.column_dimensions["E"].width = 18
ws.column_dimensions["F"].width = 18
ws.column_dimensions["G"].width = 22
ws.column_dimensions["H"].width = 18
ws.column_dimensions["I"].width = 18
ws.column_dimensions["J"].width = 80
ws.column_dimensions["K"].width = 80

wb.save(file_path)

print("\n================================")
print(f"TOPLAM İLAN: {len(all_items)}")
print("================================")
print(f"Excel oluşturuldu:")
print(file_path)